"""Adds the method slides to the deck, in place (the deck as edited in PowerPoint is kept):

  after the title            : channel charting in general
  after the scenario slide   : the EMS formulation (channel, covariance, LED); t-SNE clamped to positions
  after "Proposed: ..."      : the AoA chart; codebook and look-up table; beam choice; SNR gain

Equations and notation tables are typeset with pdflatex and placed as transparent PNGs (600 dpi) in
slides/assets/eq/; each picture keeps its LaTeX source in its alt text. Figures: make_method_figures.py.
A copy of the deck before the change goes to slides/backup/. Stops if the new slides are already there.
Slide numbers become automatic fields (they follow insertions and deletions).
Usage: python slides/add_method_slides.py [deck.pptx]
"""
import copy, os, re, shutil, subprocess, sys, tempfile, time, uuid
import fitz
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
A = os.path.join(HERE, 'assets')
DECK = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, 'CC_EMS_configuration_slides.pptx')
EQ = os.path.join(A, 'eq');  os.makedirs(EQ, exist_ok=True)
DATE = '24/10/26'
CARD, BLUE, TERRA, TEXT, MUTED, NOTE = 'EFE7DA', '1F5FA8', 'B85042', '2B2A33', '6B6670', 'F7F3EC'
rgb = lambda h: RGBColor.from_string(h)
DPI, BASE, NOTE_PT = 600, 11, 10.5                        # LaTeX base size; notation tables shown at 10.5 pt

# ====================================================================== equations (LaTeX)
EQS = {
    # channel charting in general
    'eqA': r'\begin{aligned}&\mathbf{h}_u \xrightarrow{\ \text{feature}\ } \mathbf{R}_u \xrightarrow{\ \text{LED}\ } d_{uu^\prime} '
           r'\xrightarrow{\ \text{t-SNE}\ } \mathbf{z}_u\\[3pt] &\text{goal: } d_{uu^\prime}\ \text{small} \;\Longleftrightarrow\; '
           r'\|\mathbf{z}_u-\mathbf{z}_{u^\prime}\|\ \text{small}\end{aligned}',
    # the EMS formulation (paper)
    'eqB1': r'\mathbf{h}_u(\mathcal{S}) = \mathbf{h}_u^{\mathrm{d}} + \sum_{j=1}^{2}\mathbf{H}_j^{\mathrm{o}}\,\bm{\Phi}_j\,'
            r'\mathbf{h}_{u,j}^{\mathrm{i}}, \qquad \bm{\Phi}_j = \mathrm{diag}\big(e^{j\phi_{j,1}},\dots,e^{j\phi_{j,L}}\big)',
    'eqB2': r'\mathbf{R}_u(\mathcal{S}) = \mathbb{E}\big[\,\hat{\mathbf{h}}_u(\mathcal{S})\,\hat{\mathbf{h}}_u^{H}(\mathcal{S})\,\big]',
    'eqB3': r'd_{uu^\prime}(\mathcal{S}) = \big\|\log\mathbf{R}_u(\mathcal{S}) - \log\mathbf{R}_{u^\prime}(\mathcal{S})\big\|_F',
    'eqC1': r'p_{u^\prime|u} = \frac{\exp\!\big(-d_{uu^\prime}^2/2\sigma_u^2\big)}{\sum_{w\neq u}\exp\!\big(-d_{uw}^2/2\sigma_u^2\big)},'
            r'\qquad p_{uu^\prime} = \frac{p_{u^\prime|u}+p_{u|u^\prime}}{2N_U}',
    'eqC2': r'q_{uu^\prime} = \frac{\big(1+\|\mathbf{z}_u-\mathbf{z}_{u^\prime}\|^2\big)^{-1}}{\sum_{w\neq v}\big(1+\|\mathbf{z}_w-\mathbf{z}_v\|^2\big)^{-1}}',
    'eqC3': r'\hat{\mathcal{Z}} = \operatorname*{arg\,min}_{\mathcal{Z}}\ \underbrace{\sum_{u\neq u^\prime} p_{uu^\prime}\log\frac{p_{uu^\prime}}{q_{uu^\prime}}}'
            r'_{\mathrm{KL}(P\,\|\,Q)} \quad \text{s.t. } \mathbf{z}_a = \mathbf{p}_a,\ a\in\mathcal{A}',
    'eqC4': r'\hat{\mathbf{p}}_u = \hat{\mathbf{z}}_u, \qquad \mathrm{LE}_u = \|\hat{\mathbf{z}}_u - \mathbf{p}_u\|',
    # the AoA chart (Asilomar)
    'eqD1': r'd^{(k)}_{uu^\prime} = \big\|\log\mathbf{R}_u\big(\bm{\Phi}^{(c_k)}\big) - \log\mathbf{R}_{u^\prime}\big(\bm{\Phi}^{(c_k)}\big)\big\|_F,'
            r'\qquad k = 1,\dots,K',
    'eqD2': r'\hat{\mathcal{Z}}^{(k)} = \operatorname*{arg\,min}_{\mathcal{Z}}\ \mathrm{KL}\big(P^{(k)}\,\|\,Q\big)\ \ \text{s.t. } z_a = \theta_a,\ a\in\mathcal{A},'
            r'\qquad \hat\theta^{(k)}_u = \hat z^{(k)}_u\in\mathbb{R}',
    'eqD3': r'k_u^\star = \operatorname*{arg\,max}_{k}\ \mathrm{TW}^{(k)}_u, \qquad \hat\theta_u = \hat\theta^{(k_u^\star)}_u',
    'eqD5': r'\mathrm{TW}_u = 1 - \frac{2}{\kappa(2N_U-3\kappa-1)}\sum_{u^\prime\in\mathcal{N}^{\mathrm{chart}}_\kappa(u)\setminus'
            r'\mathcal{N}^{\mathrm{LED}}_\kappa(u)}\big(r^{\mathrm{LED}}_{uu^\prime}-\kappa\big)',
    'eqD4': r'e_u = |\hat\theta_u - \theta_u|, \qquad \underset{u\in\mathcal{U}_s}{\mathrm{median}}\; e_u, \qquad '
            r'\mathrm{RMSE} = \Big(\frac{1}{|\mathcal{U}_s|}\sum_{u\in\mathcal{U}_s} e_u^2\Big)^{1/2}',
    # codebook and look-up table
    'eqE1': r'\phi^{(n)}_\ell = -\,m_\ell\,\Delta\phi_n, \qquad \Delta\phi_n = \frac{2\pi}{N_c}\Big(n-\frac{N_c}{2}\Big), \qquad n = 1,\dots,N_c',
    'eqE2': r'u_n = \big\langle \sin\theta_{\mathrm{spec}} + \Delta\phi_n/\pi \big\rangle_2',
    'eqE3': r'\Delta u = 0.886\,\frac{2}{M}, \qquad N_c = \Big\lceil\frac{2}{\Delta u}\Big\rceil_{\mathrm{even}} = 68\ (M{=}60),\ \ 272\ (M{=}240)',
    'eqE4': r'\hat c_u = \operatorname*{arg\,min}_{n}\ \big|\big\langle\sin\hat\theta_u - u_n\big\rangle_2\big|',
    # beam choice
    'eqF1': r'\hat c_u = \operatorname*{arg\,max}_{n\in\mathcal{P}_u}\ \gamma_u(n), \quad \text{sweeps} = |\mathcal{P}_u| \qquad\qquad '
            r'c^\star_u = \operatorname*{arg\,max}_{n\in\{1,\dots,N_c\}}\gamma_u(n)\ \ \text{(exhaustive)}',
    'eqF2': r'b^{(k)}_u = \operatorname*{arg\,max}_{n\in\{b^{(k-1)}_u,\ \hat c^{(k)}_u\}}\gamma_u(n), \qquad b^{(1)}_u = \hat c^{(1)}_u',
    # channel and SNR gain
    'eqG1': r'\mathbf{h}_u(n) = \mathbf{h}_u^{\mathrm{d}} + \mathbf{H}^{\mathrm{o}}\,\bm{\Phi}^{(n)}\,\mathbf{h}_u^{\mathrm{i}}',
    'eqG2': r'\gamma_u(n) = \rho\,\max_{\|\mathbf{w}\|=1}\big\|\mathbf{w}^H\mathbf{H}_u(n)\big\|^2 = \rho\,\lambda_{\max}\big(\mathbf{H}_u(n)\,\mathbf{H}_u^H(n)\big)',
    'eqG3': r'\Delta\gamma_u = 10\log_{10}\frac{\gamma_u(\hat c_u)}{\gamma_u(n_0)}\ \ \text{[dB]}, \qquad \bm{\Phi}^{(n_0)} = \mathbf{I}',
    'eqG4': r'\varepsilon_u = \frac{\big|\big\langle\sin\theta_u - u_{\hat c_u}\big\rangle_2\big|}{\Delta u}\ \ \text{[beamwidths]}',
    # the two ways to choose the training slots (update_slides_methods.py)
    'eqH1': r'c_k = \mathrm{round}\big(f_k\,N_c\big), \qquad (f_k)_{k\ge1} = \tfrac{1}{2},\ 1,\ \tfrac{1}{4},\ \tfrac{3}{4},\ '
            r'\tfrac{1}{8},\ \tfrac{3}{8},\ \tfrac{5}{8},\ \dots',
    'eqH2': r'\hat p(\theta) = \frac{1}{|\mathcal{A}|}\sum_{a\in\mathcal{A}}\delta(\theta-\theta_a)',
    'eqH3': r'c_k = \operatorname*{arg\,max}_{n\notin\{c_1,\dots,c_{k-1}\}}\ \big|\{a\in\mathcal{A}_k:\ |\sin\theta_a - u_n| \le r\}\big|',
}
# notation tables: name -> (box width [in], rows of (symbol, meaning))
NOTES = {
    'notA': (5.35, [(r'u', r'a user (test point), at position $\mathbf{p}_u$'),
                    (r'\mathbf{h}_u\in\mathbb{C}^{N_{\mathrm{BS}}}', r'its uplink channel at the BS ($N_{\mathrm{BS}}$ antennas)'),
                    (r'\mathbf{R}_u', r'covariance of $\mathbf{h}_u$ (the feature)'),
                    (r'd_{uu^\prime}', r'LED: dissimilarity of users $u$ and $u^\prime$ (next slides)'),
                    (r'\mathbf{z}_u\in\mathbb{R}^{D^\prime}', r'place of $u$ in the chart ($D^\prime = 1$ or $2$)')]),
    'notB1': (4.8, [(r'\mathbf{h}_u^{\mathrm{d}}', r'direct UE--BS channel ($\mathbf{0}$ for full-NLOS users)'),
                    (r'\mathbf{h}^{\mathrm{i}}_{u,j}\in\mathbb{C}^{L}', r'UE $\to$ EMS $j$ (incident)'),
                    (r'\mathbf{H}^{\mathrm{o}}_j\in\mathbb{C}^{N_{\mathrm{BS}}\times L}', r'EMS $j\to$ BS (outgoing)'),
                    (r'\phi_{j,\ell}', r'phase of element $\ell$ of EMS $j$; $L = M\times M$ elements'),
                    (r'\mathcal{S}', r'configuration $\{\bm{\Phi}_1,\bm{\Phi}_2\}$'),
                    (r'\hat{\mathbf{h}}_u,\ \mathbb{E}', r'noisy channel estimate; mean over $T = 500$ pilots'),
                    (r'\log,\ \|\cdot\|_F', r'matrix logarithm; Frobenius norm')]),
    'notB2': (7.25, [(r'\sigma_u', r'kernel width: each user gets about \emph{perplexity} (150) effective neighbours'),
                     (r'N_U,\ \mathbf{z}_u\in\mathbb{R}^2', r'number of users; position of user $u$ in the chart'),
                     (r'\mathcal{A},\ \mathbf{p}_a', r'anchors (15\,\% of the users) and their known positions'),
                     (r'\mathrm{KL}', r'Kullback--Leibler divergence, minimised by gradient descent; anchors reset at every iteration'),
                     (r'\mathrm{LE}_u', r'localisation error of user $u$ (true position $\mathbf{p}_u$)')]),
    'notD': (4.95, [(r'\theta_u', r'azimuth of user $u$ seen from the EMS (angle of arrival)'),
                    (r'c_k', r'codeword of training slot $k$ (uninformed or informed: previous slide)'),
                    (r'\bm{\Phi}^{(c)}', r'phase profile of codeword $c$ (next slide)'),
                    (r'\mathcal{N}_\kappa(u)', r'$\kappa$ nearest users of $u$ (chart or LED); $\kappa$ = 5\,\% of users'),
                    (r'r^{\mathrm{LED}}_{uu^\prime}', r'rank of $u^\prime$ among the LED neighbours of $u$'),
                    (r'\mathcal{U}_s', r'scored users: reached by the EMS, not anchors')]),
    'notE': (3.75, [(r'm_\ell', r'column of element $\ell$ ($1,\dots,M$): ramp along x only'),
                    (r'\theta_{\mathrm{spec}}', r'incidence that a zero-phase EMS (a mirror) reflects to the BS'),
                    (r'u_n', r'$\sin$ of the incidence that codeword $n$ reflects to the BS'),
                    (r'\langle x\rangle_2', r'$x$ wrapped into $[-1,1)$: a $\lambda/2$ array repeats every 2 in $\sin\theta$'),
                    (r'\Delta u', r'half-power beamwidth in $\sin\theta$ ($1.7^\circ$ at broadside, $M = 60$)'),
                    (r'N_c', r'number of codewords; $\lceil\cdot\rceil_{\mathrm{even}}$: rounded up to even')]),
    'notF': (3.7, [(r'\gamma_u(n)', r'SNR of $u$ when the EMS applies codeword $n$'),
                   (r'\mathcal{P}_u', r'codewords measured for user $u$, one sweep each'),
                   (r'\hat c^{(k)}_u', r'look-up proposal after training slot $k$'),
                   (r'b^{(k)}_u', r'beam kept after slot $k$ (t-SNE + verify)')]),
    'notG': (4.95, [(r'\rho', r'$P_{\mathrm{tx}}/\sigma_n^2$ (23 dBm, $-92$ dBm noise)'),
                    (r'\mathbf{H}_u(n)', r'channel of $u$ at the BS for codeword $n$ (the dataset keeps 2 channel columns per user; '
                                         r'one column gives $\rho\|\mathbf{h}_u(n)\|^2$)'),
                    (r'\mathbf{w},\ \lambda_{\max}', r'BS combiner ($N_{\mathrm{BS}} = 32$); largest eigenvalue'),
                    (r'n_0', r'zero-phase codeword ($n_0 = N_c/2$): the EMS acts as a mirror'),
                    (r'\hat c_u,\ \mathcal{U}_s', r'chosen codeword; scored users (previous slides)')]),
    'notH': (5.3, [(r'f_k', r'fraction of the codebook for slot $k$: the gaps are halved level by level'),
                   (r'N_c,\ u_n', r'number of codewords; direction ($\sin\theta$) served by codeword $n$'),
                   (r'\hat p(\theta)', r'prior PDF of the user angles at the EMS, from the anchors'),
                   (r'\mathcal{A},\ \theta_a', r'anchors (15\,\% of the users) and their azimuths'),
                   (r'\mathcal{A}_k', r'anchors not yet covered by slots $1,\dots,k-1$'),
                   (r'r', r'coverage radius in $\sin\theta$ ($0.15$), halved once (almost) every anchor is covered')]),
}
NOTE_PAD = 0.15                                            # inner margin of the notation box [in]

PRE = r'''\documentclass[11pt,border=1.5pt,multi={eqn,tab}]{standalone}
\usepackage{amsmath,amssymb,bm,xcolor,array,tabularx}
\newenvironment{eqn}{\color[HTML]{2B2A33}$\displaystyle}{$}
\newenvironment{tab}{\color[HTML]{2B2A33}\renewcommand{\arraystretch}{1.18}}{}
\begin{document}
'''
SIZE = {}                                                  # name -> (width, height) [in] at BASE pt


def typeset():
    items = [(n, r'\begin{eqn}' + b + r'\end{eqn}') for n, b in EQS.items()]
    for n, (w, rows) in NOTES.items():
        wl = (w - 2 * NOTE_PAD) * BASE / NOTE_PT
        body = r' \\ '.join(f'{s} & {t}' for s, t in rows)
        items.append((n, r'\begin{tab}\begin{tabularx}{' + f'{wl:.3f}in' + r'}{@{}>{$}l<{$}@{\hspace{7pt}}X@{}}' + body
                      + r'\end{tabularx}\end{tab}'))
    d = tempfile.mkdtemp()
    with open(os.path.join(d, 'eq.tex'), 'w', encoding='utf-8') as f:
        f.write(PRE + '\n'.join(b for _, b in items) + '\n\\end{document}\n')
    r = subprocess.run(['pdflatex', '-interaction=nonstopmode', '-halt-on-error', 'eq.tex'], cwd=d,
                       capture_output=True, text=True, timeout=300)
    if r.returncode:
        print(r.stdout[-4000:]);  raise SystemExit('pdflatex failed')
    doc = fitz.open(os.path.join(d, 'eq.pdf'))
    assert doc.page_count == len(items), (doc.page_count, len(items))
    for (n, _), page in zip(items, doc):
        page.get_pixmap(dpi=DPI, alpha=True).save(os.path.join(EQ, n + '.png'))
        SIZE[n] = (page.rect.width / 72, page.rect.height / 72)


# ====================================================================== drawing helpers (as build_slides.py)
def runs(p, text, size, color, bold=False, italic=False, font='Calibri'):
    for k, part in enumerate(re.split(r'\*\*', text)):
        if not part:
            continue
        r = p.add_run();  r.text = part
        f = r.font;  f.size = Pt(size);  f.bold = bold or (k % 2 == 1);  f.italic = italic
        f.color.rgb = rgb(color);  f.name = font


def textbox(s, x, y, w, h, paras, size=12.5, color=TEXT, bold=False, italic=False, align=PP_ALIGN.LEFT, space=4):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame;  tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, t in enumerate(paras if isinstance(paras, list) else [paras]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align;  p.space_after = Pt(space)
        runs(p, t, size, color, bold, italic)
    return tb


def card(s, x, y, w, h, header=None, body=(), hcolor=BLUE, size=14):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.10
    sh.fill.solid();  sh.fill.fore_color.rgb = rgb(CARD);  sh.line.fill.background();  sh.shadow.inherit = False
    top = y + 0.1
    if header:
        textbox(s, x + 0.2, top, w - 0.4, 0.38, header, size=17, color=hcolor, bold=True);  top += 0.47
    if body:
        textbox(s, x + 0.2, top, w - 0.4, h - (top - y) - 0.1, list(body), size=size)
    return sh


def figure(s, path, x, y, w, h, align='center'):
    """Picture fitted into the box (aspect kept); returns its (x, y, w, h)."""
    from PIL import Image
    iw, ih = Image.open(path).size
    sc = min(w / iw, h / ih);  pw, ph = iw * sc, ih * sc
    px = x + (w - pw) / 2 if align == 'center' else x
    s.shapes.add_picture(path, Inches(px), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))
    return px, y + (h - ph) / 2, pw, ph


def eq(s, name, x, y, pt=16, maxw=None):
    """Equation image at pt points (shrunk to maxw if wider); returns (width, height) [in]."""
    w, h = SIZE[name];  sc = pt / BASE
    if maxw and w * sc > maxw:
        sc = maxw / w
    pic = s.shapes.add_picture(os.path.join(EQ, name + '.png'), Inches(x), Inches(y), Inches(w * sc), Inches(h * sc))
    pic._element.nvPicPr.cNvPr.set('descr', EQS.get(name, name))
    return w * sc, h * sc


def notation(s, name, x, y):
    w, rows = NOTES[name];  iw, ih = SIZE[name];  sc = NOTE_PT / BASE
    h = 0.42 + ih * sc + NOTE_PAD
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.06
    sh.fill.solid();  sh.fill.fore_color.rgb = rgb(NOTE);  sh.line.color.rgb = rgb('D9D0C1');  sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    textbox(s, x + NOTE_PAD, y + 0.08, w - 2 * NOTE_PAD, 0.3, 'Notation', size=12, color=MUTED, bold=True)
    pic = s.shapes.add_picture(os.path.join(EQ, name + '.png'), Inches(x + NOTE_PAD), Inches(y + 0.42), Inches(iw * sc), Inches(ih * sc))
    pic._element.nvPicPr.cNvPr.set('descr', 'notation: ' + '; '.join(f'{a}: {b}' for a, b in rows))
    return y + h


def badge(s, x, y, num):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.3), Inches(0.3))
    sh.fill.solid();  sh.fill.fore_color.rgb = rgb(BLUE);  sh.line.fill.background();  sh.shadow.inherit = False
    tf = sh.text_frame;  tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0];  p.alignment = PP_ALIGN.CENTER;  runs(p, str(num), 12, 'FFFFFF', bold=True)


def step(s, x, y, w, num, head, eqs, note=None, pt=16, hgap=0.42, gap=0.14):
    """Numbered step: header, its equations, an optional note; returns the y below it."""
    if num is not None:
        badge(s, x, y + 0.01, num)
    hx = x + (0.42 if num is not None else 0)
    textbox(s, hx, y, w - (hx - x), 0.32, head, size=14.5, bold=True)
    yy = y + hgap
    for e in eqs:
        p = pt[e] if isinstance(pt, dict) else pt
        _, eh = eq(s, e, hx + 0.05, yy, p, maxw=w - (hx - x) - 0.05);  yy += eh + 0.08
    if note:
        textbox(s, hx + 0.05, yy, w - (hx - x) - 0.05, 0.3, note, size=12, color=MUTED, italic=True);  yy += 0.34
    return yy + gap


# ====================================================================== deck
def footer(prs, s):
    for ph in prs.slide_layouts[1].placeholders:
        idx = ph.placeholder_format.idx
        if idx in (12, 13, 14):
            el = copy.deepcopy(ph._element);  s.shapes._spTree.append(el)
            s.shapes[-1].text_frame.text = {12: 'CONFIDENTIAL', 13: '0', 14: DATE}[idx]


def new_slide(prs, title, subtitle=None):
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = title
    for r in s.shapes.title.text_frame.paragraphs[0].runs:
        r.font.size = Pt(28)
    for ph in list(s.placeholders):                        # the empty body placeholder is not used
        if ph.placeholder_format.idx == 1:
            ph._element.getparent().remove(ph._element)
    if subtitle:
        textbox(s, 0.1, 0.7, 12.0, 0.36, subtitle, size=15, color=TERRA, italic=True)
    footer(prs, s)
    return s


def slide_number_fields(prs):
    """The slide-number placeholder of every slide becomes a field: PowerPoint keeps it right."""
    for i, s in enumerate(prs.slides, 1):
        for ph in s.placeholders:
            if ph.placeholder_format.idx != 13:
                continue
            p = ph.text_frame.paragraphs[0]._p
            r0 = p.find(qn('a:r'));  r0 = r0 if r0 is not None else p.find(qn('a:fld'))
            rPr = copy.deepcopy(r0.find(qn('a:rPr'))) if r0 is not None and r0.find(qn('a:rPr')) is not None else etree.Element(qn('a:rPr'))
            for e in p.findall(qn('a:r')) + p.findall(qn('a:fld')):
                p.remove(e)
            fld = etree.Element(qn('a:fld'));  fld.set('id', '{%s}' % str(uuid.uuid4()).upper());  fld.set('type', 'slidenum')
            fld.append(rPr);  t = etree.SubElement(fld, qn('a:t'));  t.text = str(i)
            end = p.find(qn('a:endParaRPr'))
            if end is not None:
                end.addprevious(fld)
            else:
                p.append(fld)


def title_of(s):
    t = s.shapes.title
    return t.text_frame.text if t is not None else ''


paper = lambda n: os.path.join(A, 'paper', n + '.png')
meth = lambda n: os.path.join(A, 'method', n + '.png')


def build_A(prs):
    s = new_slide(prs, 'Channel charting: a map of the users learned from their channels',
                  'Unsupervised: users with similar channels end up close together, with no position needed.')
    figure(s, meth('cc_geometries'), 0.3, 1.12, 7.0, 5.5)
    textbox(s, 0.45, 6.66, 6.9, 0.3, 'Toy example made for these slides (30 GHz, 8-antenna BS, line of sight + two wall reflections, '
            '500 users), drawn as Fig. 3 of Studer et al., IEEE Access 2018.', size=10, color=MUTED, italic=True)
    card(s, 7.55, 1.15, 5.35, 2.95, 'How it works', [
        '**1  Channel.** Each user sends an uplink pilot; the BS estimates its channel vector.',
        '**2  Feature.** The raw channel cannot be compared: its phase turns every centimetre at 30 GHz. '
        'Its covariance changes slowly with position; users are compared by a distance between covariances.',
        '**3  Chart.** A dimensionality reduction (here t-SNE) places the users in 1-D or 2-D so that similar '
        'channels are neighbours.'], size=13)
    eq(s, 'eqA', 7.75, 4.25, pt=15, maxw=4.95)
    notation(s, 'notA', 7.55, 5.2)


def build_B1(prs):
    s = new_slide(prs, 'Channel charting with EMSs: channel, feature, dissimilarity',
                  'The EMS configuration enters every channel, hence every feature and the chart itself.')
    y = step(s, 0.45, 1.25, 7.4, 1, 'Channel at the BS: the direct path plus one path per EMS', ['eqB1'],
             'Paths from ray tracing (Sionna); the EMS phases are what we control.', pt=15)
    y = step(s, 0.45, y, 7.4, 2, 'Feature: covariance over repeated noisy pilots', ['eqB2'],
             'Changes slowly with position; insensitive to the carrier phase.', pt=16)
    y = step(s, 0.45, y, 7.4, 3, 'Dissimilarity: log-Euclidean distance (LED)', ['eqB3'],
             'Compares the received power (log scale) and the spatial signature of two users.', pt=16)
    card(s, 0.45, y + 0.05, 7.4, 0.8, None, [
        '**Our TWC paper:** pick the EMS phases from a codebook to get the best chart (one static codeword), '
        'or fuse the charts of several slots, one codeword each.'], size=13)
    fx, fy, fw, fh = figure(s, meth('sre_features'), 8.1, 1.12, 4.8, 3.15)
    textbox(s, 8.1, fy + fh + 0.02, 4.8, 0.3, 'Toy example: A and B (1.5 m apart) receive power from the same directions; C does not.',
            size=10, color=MUTED, italic=True)
    notation(s, 'notB1', 8.1, fy + fh + 0.5)


def build_B2(prs):
    s = new_slide(prs, 't-SNE turns dissimilarities into a chart; anchors pin it to positions',
                  'Semi-supervised t-SNE: 15% of the users (anchors) have a known position and stay fixed.')
    y = step(s, 0.45, 1.15, 7.3, None, 'Neighbours in the channel space (Gaussian kernel)', ['eqC1'], pt=13, hgap=0.36, gap=0.04)
    y = step(s, 0.45, y, 7.3, None, 'Neighbours in the chart (Student-t kernel: heavy tail)', ['eqC2'], pt=13, hgap=0.36, gap=0.04)
    y = step(s, 0.45, y, 7.3, None, 'Match the two with anchors fixed; the chart gives the position', ['eqC3', 'eqC4'],
             pt=13, hgap=0.36, gap=0.04)
    notation(s, 'notB2', 0.45, y + 0.04)
    for k, (n, cap) in enumerate((('tp_ground_truth', 'true positions'), ('tp_no_ems', 'chart, no EMS'),
                                  ('tp_ems_spec', 'chart, specular EMS'), ('tp_ems_best', 'chart, optimised EMS'))):
        x = 7.95 + (k % 2) * 2.5;  yy = 1.2 + (k // 2) * 2.3
        textbox(s, x, yy, 2.4, 0.28, cap, size=11.5, bold=True, align=PP_ALIGN.CENTER)
        figure(s, paper(n), x, yy + 0.3, 2.4, 1.85)
    textbox(s, 7.95, 5.85, 4.9, 0.6, 'From our TWC paper: S-tSNE charts of the test points, colour = true y coordinate; '
            'the chart is already in metres thanks to the anchors.', size=10.5, color=MUTED, italic=True)


def build_D(prs):
    s = new_slide(prs, 'AoA chart: the same formulation, with the angle as label',
                  'In this paper, per EMS: K training slots, one 1-D chart per slot, anchors labelled with their azimuth.')
    y = step(s, 0.45, 1.2, 7.25, 1, 'Slot k: one codeword for all users, as in our TWC paper', ['eqD1'], pt=14)
    y = step(s, 0.45, y - 0.04, 7.25, 2, '1-D t-SNE (same p and q), anchors clamped to their azimuth', ['eqD2'], pt=14)
    y = step(s, 0.45, y - 0.04, 7.25, 3, 'Per user, the most trustworthy of the K charts (no labels)', ['eqD3', 'eqD5'],
             pt={'eqD3': 14, 'eqD5': 12.5})
    y = step(s, 0.45, y - 0.04, 7.25, 4, 'Angle error over the scored users', ['eqD4'],
             'Median: the typical user. RMSE: dominated by the few users placed in a wrong group.', pt=14)
    fx, fy, fw, fh = figure(s, meth('aoa_chart_1d'), 7.95, 1.12, 4.95, 3.65)
    notation(s, 'notD', 7.95, fy + fh + 0.08)


def build_E(prs):
    s = new_slide(prs, 'From angle to codeword: steering codebook and look-up table',
                  'Each codeword is a linear phase ramp across the EMS; the look-up table picks the beam nearest to the estimate.')
    figure(s, meth('codebook_beams'), 0.35, 1.12, 8.75, 3.3)
    notation(s, 'notE', 9.2, 1.15)
    y1 = step(s, 0.45, 4.6, 6.2, 1, 'Codeword n: a linear phase ramp along x', ['eqE1'], pt=14)
    step(s, 0.45, y1 - 0.04, 6.2, 2, 'Beamwidth and number of codewords', ['eqE3'], pt=14)
    y2 = step(s, 6.85, 4.6, 6.05, 3, 'Steering law (λ/2 spacing): incidence served by n', ['eqE2'], pt=15)
    step(s, 6.85, y2 - 0.04, 6.05, 4, 'Look-up table: the beam nearest to the estimate', ['eqE4'], pt=15)


def build_F(prs):
    s = new_slide(prs, 'Choosing the beam: look-up, verify, search',
                  'Every scheme measures some codewords and keeps the strongest; they differ in which ones and how many.')
    eq(s, 'eqF1', 0.5, 1.22, pt=14, maxw=8.4)
    figure(s, meth('beam_choice'), 0.35, 1.95, 8.6, 3.95)
    textbox(s, 0.5, 5.95, 3.1, 0.3, '**t-SNE + verify:**', size=13)
    eq(s, 'eqF2', 2.2, 5.9, pt=13.5, maxw=6.6)
    textbox(s, 0.5, 6.62, 8.4, 0.3, 'Illustration: one user, single-path beam gain, made-up anchor groups; '
            'performance: see the results.', size=10, color=MUTED, italic=True)
    card(s, 9.2, 1.15, 3.7, 4.05, 'The schemes', [
        '**Exhaustive:** all codewords (68 to 272 sweeps): the upper bound.',
        '**Greedy:** one codeword per sweep along a fixed list (with the prior: the beams at the user groups first).',
        '**Hierarchical:** 4 coarse probes, then ±δ around the best so far, δ halved.',
        '**t-SNE + look-up:** K training slots, shared by all users; the proposal is not measured.',
        '**t-SNE + R neighbours:** the proposal and its R nearest codewords measured (K + R + 1 sweeps).',
        '**+ verify:** each proposal measured, the stronger kept: SNR never drops (2K sweeps).',
        '**+ neighbour search:** ±1, ±2, ... around the kept beam.'], size=11.5)
    notation(s, 'notF', 9.2, 5.3)


def build_G(prs):
    s = new_slide(prs, 'What the chosen codeword does to the channel, and the SNR gain',
                  'SNR at the BS with the best combiner; the gain is measured against the unsteered EMS (zero phase).')
    y = step(s, 0.45, 1.2, 7.2, 1, 'The codeword changes only the EMS path', ['eqG1'],
             'For full-NLOS users the direct channel is zero: the EMS path is the whole channel.', pt=16)
    y = step(s, 0.45, y - 0.04, 7.2, 2, 'SNR with the best BS combiner', ['eqG2'], pt=15)
    y = step(s, 0.45, y - 0.04, 7.2, 3, 'SNR gain of the chosen beam over the zero-phase EMS', ['eqG3'],
             'A zero-phase EMS is a mirror: it serves the specular direction only, hence gains of tens of dB.', pt=15)
    y = step(s, 0.45, y - 0.04, 7.2, 4, 'Pointing error of the chosen beam', ['eqG4'],
             'Reported over the scored users: median SNR gain; median and RMSE of the pointing error.', pt=15)
    figure(s, meth('gain_vs_pointing'), 7.95, 1.12, 4.95, 3.2)
    notation(s, 'notG', 7.95, 4.5)


NEW = [('after_title', build_A), ('after_scenario', build_B1), ('after_scenario', build_B2),
       ('after_proposed', build_D), ('after_proposed', build_E), ('after_proposed', build_F), ('after_proposed', build_G)]
ANCHOR = {'after_scenario': 'Smart radio environment', 'after_proposed': 'Proposed: channel charting of the angle'}

if __name__ == '__main__':
    prs = Presentation(DECK)
    titles = [title_of(s) for s in prs.slides]
    new_titles = ['Channel charting: a map of the users', 'Channel charting with EMSs', 't-SNE turns dissimilarities',
                  'AoA chart: the same formulation', 'From angle to codeword', 'Choosing the beam', 'What the chosen codeword']
    if any(t.startswith(n) for t in titles for n in new_titles):
        raise SystemExit('method slides already in the deck: nothing done')
    typeset()
    os.makedirs(os.path.join(os.path.dirname(DECK), 'backup'), exist_ok=True)
    bk = os.path.join(os.path.dirname(DECK), 'backup', time.strftime('CC_EMS_configuration_slides_%Y%m%d_%H%M%S.pptx'))
    shutil.copy2(DECK, bk)
    lst = prs.slides._sldIdLst
    old = list(lst)
    built = []
    for where, fn in NEW:
        fn(prs);  built.append((where, list(lst)[-1]))
    pos = {'after_title': 0}
    for key, start in ANCHOR.items():
        hit = [i for i, t in enumerate(titles) if t.startswith(start)]
        assert len(hit) == 1, (start, hit)
        pos[key] = hit[0]
    order = []
    for i, e in enumerate(old):
        order.append(e)
        order += [el for where, el in built if pos[where] == i]
    for e in list(lst):
        lst.remove(e)
    for e in order:
        lst.append(e)
    slide_number_fields(prs)
    tmp = DECK + '.tmp.pptx'
    prs.save(tmp)
    try:
        os.replace(tmp, DECK)
    except PermissionError:
        raise SystemExit(f'the deck is open in PowerPoint: close it and run again (new version left in {tmp})')
    print(f'{len(prs.slides)} slides; backup {bk}')
