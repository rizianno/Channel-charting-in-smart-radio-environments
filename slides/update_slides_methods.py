"""Second update of the deck, in place (the deck as edited in PowerPoint is kept):

  - new slide after "Proposed: ...": in this paper, two ways to choose the training slots (uninformed,
    informed with an angle prior), with the real EMS2 prior and slots (make_method_figures.py)
  - Case 1 / Case 2 "installed" slides: the figure box gets the per-EMS CDF of the AoA error, every K,
    uninformed above informed (experiments/plot_aoa_cdf.py, 60x60)
  - wording: "our TWC paper" for the earlier work, "in this paper" for this one; no coarse-to-fine,
    random, clusters, synthetic or "real" 60x60
  - the AoA-chart notation and the beam-choice figure are redrawn with that wording
Equations and notation come from add_method_slides.py (EQS, NOTES). A copy of the deck before the
change goes to slides/backup/. Stops if the new slide is already there.
Usage: python slides/update_slides_methods.py [deck.pptx]
"""
import os, shutil, sys, time
from pptx import Presentation
from pptx.util import Inches, Pt
import add_method_slides as am

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
DECK = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, 'CC_EMS_configuration_slides.pptx')
RES = os.path.join(REPO, 'results_aoa', 'serving_deployments', 'est_G60dB', 'figures', 'aoa_cdf')
ASSET = os.path.join(HERE, 'assets', 'results');  os.makedirs(ASSET, exist_ok=True)

# (slide title start, paragraph text, new text with **bold**, or None to delete the paragraph)
EDITS = [
    ('Channel charting with EMSs', 'Paper: pick the EMS phases from a codebook to get the best chart (one static codeword), '
     'or fuse the charts of several slots, one codeword each.',
     '**Our TWC paper:** pick the EMS phases from a codebook to get the best chart (one static codeword), '
     'or fuse the charts of several slots, one codeword each.'),
    ('t-SNE turns dissimilarities', 'From the paper: S-tSNE charts of the test points, colour = true y coordinate; '
     'the chart is already in metres thanks to the anchors.',
     'From our TWC paper: S-tSNE charts of the test points, colour = true y coordinate; '
     'the chart is already in metres thanks to the anchors.'),
    ('Paper result:', 'Paper result: EMS configurations make the chart more accurate',
     'Our TWC paper: EMS configurations improve the chart'),
    ('Proposed: channel charting', "The paper's static slots become training acquisitions; the chart learns the AoA instead of the position.",
     'In this paper, the static slots of our TWC paper become training acquisitions; the chart learns the AoA instead of the position.'),
    ('Proposed: channel charting', 'Static EMS configurations, as in the paper. Informed: chosen from the angular spread of the anchors.',
     'Static EMS configurations, as in our TWC paper: **uninformed** or **informed** (next slide).'),
    ('AoA chart:', 'Slot k: one codeword for all users (the paper’s static slots)',
     'Slot k: one codeword for all users, as in our TWC paper'),
    ('AoA chart:', 'Per EMS: K training slots, one 1-D chart per slot, anchors labelled with their azimuth at the EMS.',
     'In this paper, per EMS: K training slots, one 1-D chart per slot, anchors labelled with their azimuth.'),
    ('Choosing the beam', 'Greedy: one codeword per sweep along a fixed list (clusters first: beams at the anchors).',
     '**Greedy:** one codeword per sweep along a fixed list (with the prior: the beams at the user groups first).'),
    ('Evaluation setup', 'Real 60x60 EMS (ray-traced); synthetic 120x120 ... 240x240 (fitted spherical-wave model).',
     'EMS sizes from 60x60 to 240x240 elements.'),
    ('Evaluation setup', 'Informed: training beams cover the angular groups of the anchors (prior).',
     '**Uninformed:** the slots spread over the whole codebook (no prior).'),
    ('Evaluation setup', 'Coarse-to-fine: one fixed ever-finer order (no prior).',
     '**Informed:** the slots placed by the angle prior of the anchors.'),
    ('Evaluation setup', 'Random: random codewords (no prior), 10 draws.', None),
    ('Evaluation setup', 'Greedy search and hierarchical search over the beams pointing at the user groups (same prior).',
     '**Greedy search** and **hierarchical search**; with the prior, over the beams pointing at the user groups.'),
    ('Case 1: results', 'Key numbers (real 60x60)', 'Key numbers (60x60)'),
    ('Case 1: results', 't-SNE, t-SNE + neighbours, greedy and hierarchical search (clusters), exhaustive.',
     't-SNE, t-SNE + neighbours, greedy and hierarchical search, exhaustive.'),
    ('Case 2: results', 'Real 60x60: a guided search is as good early on; on larger surfaces the chart leads.',
     '60x60: a guided search is as good early on; on larger surfaces the chart leads.'),
    ('Case 2: results', 'Median SNR gain vs sweeps, EMS2 trained alone, real 60x60.', 'Median SNR gain vs sweeps, EMS2, 60x60.'),
    ('Training choice matters', 'Placing the training beams on the user groups (informed) beats random or fixed training.',
     'Placing the training slots on the user groups (informed) beats the uninformed order.'),
    ('Training choice matters', 'Median angle error vs K: informed vs coarse-to-fine vs random, EMS1 and EMS2 (trained alone).',
     'Median angle error vs K: informed vs uninformed, EMS1 and EMS2.'),
    ('Training choice matters', 'Numbers exist (EMS2, real 60x60): K = 3: 3.7 deg informed vs 37 deg random; K = 10: 1.9 vs 14 deg.',
     'EMS2, 60x60: K = 3: 3.7 deg informed vs 42.9 deg uninformed; K = 6: 3.1 vs 3.5 deg.'),
    ('Training choice matters', 'Informed training needs one beam per user group: about 3 for EMS2, instead of 6 (coarse-to-fine) or 10+ (random).',
     'Informed training needs **one beam per user group**: about 3 for EMS2, instead of about 6 without the prior.'),
]


def find(prs, start):
    hit = [s for s in prs.slides if am.title_of(s).startswith(start)]
    assert len(hit) == 1, (start, len(hit))
    return hit[0]


def set_para(para, text):
    """Replace a paragraph's text. Plain text into one run: that run keeps its format (titles keep the
    theme); otherwise rebuilt with the size, colour and italics of the first run, **bold** kept."""
    if '**' not in text and len(para.runs) == 1:
        para.runs[0].text = text
        return
    r0 = para.runs[0].font
    size, italic = (r0.size.pt if r0.size else 14), r0.italic
    try:
        color = str(r0.color.rgb)
    except AttributeError:
        color = am.TEXT
    for r in list(para.runs):
        para._p.remove(r._r)
    am.runs(para, text, size, color, italic=italic)


def edit_text(prs):
    for start, old, new in EDITS:
        s = find(prs, start);  done = 0
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            for para in list(sh.text_frame.paragraphs):
                if para.text == old:
                    if new is None:
                        para._p.getparent().remove(para._p)
                    else:
                        set_para(para, new)
                    done += 1
        assert done == 1, (start, old, done)


def replace_picture(s, pic, path, keep='width'):
    """New image at the same place (same width, or same height)."""
    from PIL import Image
    iw, ih = Image.open(path).size
    x, y, w, h = pic.left, pic.top, pic.width, pic.height
    if keep == 'width':
        h = int(w * ih / iw)
    else:
        w = int(h * iw / ih)
    descr = pic._element.nvPicPr.cNvPr.get('descr')
    pic._element.getparent().remove(pic._element)
    new = s.shapes.add_picture(path, x, y, w, h)
    if descr:
        new._element.nvPicPr.cNvPr.set('descr', descr)


def redraw(prs):
    s = find(prs, 'AoA chart:')                           # notation with the new training wording
    pic = [p for p in s.shapes if p.shape_type == 13 and (p._element.nvPicPr.cNvPr.get('descr') or '').startswith('notation:')]
    assert len(pic) == 1
    w, rows = am.NOTES['notD']
    replace_picture(s, pic[0], os.path.join(am.EQ, 'notD.png'))
    pic = [p for p in s.shapes if p.shape_type == 13 and (p._element.nvPicPr.cNvPr.get('descr') or '').startswith('notation:')]
    pic[0]._element.nvPicPr.cNvPr.set('descr', 'notation: ' + '; '.join(f'{a}: {b}' for a, b in rows))
    s = find(prs, 'Choosing the beam')                    # beam-choice figure: the largest picture
    pic = max((p for p in s.shapes if p.shape_type == 13), key=lambda p: p.width * p.height)
    replace_picture(s, pic, am.meth('beam_choice'))


def fill_cases(prs):
    for start, view in (('Case 1: EMS1 installed', 'EMS1'), ('Case 2: EMS2 installed', 'EMS2')):
        s = find(prs, start)
        box = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith('FIGURE TO DECIDE')]
        assert len(box) == 1, start
        x, y, w, h = (v / 914400 for v in (box[0].left, box[0].top, box[0].width, box[0].height))
        box[0]._element.getparent().remove(box[0]._element)
        src = os.path.join(RES, f'cdf_aoa_{view}_M60_raytraced_v.png')
        dst = os.path.join(ASSET, f'cdf_aoa_{view}_60x60.png')
        shutil.copy2(src, dst)
        am.figure(s, dst, x, y, w, h)


def build_H(prs):
    s = am.new_slide(prs, 'In this paper: two ways to choose the training slots',
                     'Uninformed: nothing is known about the users. Informed: a prior PDF of the user angles puts the slots where the users are.')
    y = am.step(s, 0.45, 1.25, 6.95, 1, 'Uninformed: spread the slots over the whole codebook', ['eqH1'],
                'Each new slot halves the largest gap between the slots taken so far.', pt=15)
    y = am.step(s, 0.45, y, 6.95, 2, 'Informed: a prior PDF of the user angles at the EMS', ['eqH2'],
                'Learned offline: here from the anchors (in practice, earlier CSI or traffic statistics).', pt=15)
    y = am.step(s, 0.45, y, 6.95, 3, 'Each slot covers most of the prior not yet covered', ['eqH3'],
                'Slots land on the user groups; none is spent where there are no users.', pt=15)
    am.card(s, 0.45, y + 0.02, 6.95, 0.8, None, [
        '**Same prior for the baselines:** with it, greedy and hierarchical search start from the beams '
        'pointing at the user groups; without it, they search the whole codebook.'], size=13)
    fx, fy, fw, fh = am.figure(s, am.meth('training_slots'), 7.6, 1.15, 5.3, 3.55)
    am.notation(s, 'notH', 7.6, fy + fh + 0.1)
    return s


if __name__ == '__main__':
    prs = Presentation(DECK)
    if any(am.title_of(s).startswith('In this paper: two ways') for s in prs.slides):
        raise SystemExit('the two-methods slide is already in the deck: nothing done')
    am.typeset()
    os.makedirs(os.path.join(os.path.dirname(DECK), 'backup'), exist_ok=True)
    bk = os.path.join(os.path.dirname(DECK), 'backup', time.strftime('CC_EMS_configuration_slides_%Y%m%d_%H%M%S.pptx'))
    shutil.copy2(DECK, bk)
    edit_text(prs)
    redraw(prs)
    fill_cases(prs)
    lst = prs.slides._sldIdLst
    titles = [am.title_of(s) for s in prs.slides]
    anchor = list(lst)[[i for i, t in enumerate(titles) if t.startswith('Proposed: channel charting')][0]]
    build_H(prs)
    new = list(lst)[-1];  lst.remove(new);  anchor.addnext(new)
    am.slide_number_fields(prs)
    tmp = DECK + '.tmp.pptx'
    prs.save(tmp)
    try:
        os.replace(tmp, DECK)
    except PermissionError:
        raise SystemExit(f'the deck is open in PowerPoint: close it and run again (new version left in {tmp})')
    print(f'{len(prs.slides)} slides; backup {bk}')
