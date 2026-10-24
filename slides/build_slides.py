"""Builds the meeting slides from the group template (python-pptx).

Template: 'permittivity_slides_meeting_template 1.pptx' (Desktop): its master (logos, footer bar)
and styles are kept; its own slides are removed. Slides are drawn in the template's style: title
placeholder, italic subtitle line, beige rounded cards with a bold header, figures; figure ideas not
yet decided are dashed boxes with the idea written in them.
Output: slides/CC_EMS_configuration_slides.pptx
Usage: python slides/build_slides.py
"""
import copy, os, re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
TEMPLATE = r'C:\Users\rezaa\OneDrive - Politecnico di Milano\Desktop\permittivity_slides_meeting_template 1.pptx'
OUT = os.path.join(HERE, 'CC_EMS_configuration_slides.pptx')
A = os.path.join(HERE, 'assets')
DATE = '24/10/26'
CARD, BLUE, TERRA, TEXT, MUTED = 'EFE7DA', '1F5FA8', 'B85042', '2B2A33', '6B6670'
rgb = lambda h: RGBColor.from_string(h)

prs = Presentation(TEMPLATE)
for sid in list(prs.slides._sldIdLst):                     # drop the template's own slides
    prs.part.drop_rel(sid.rId);  prs.slides._sldIdLst.remove(sid)
L_TITLE, L_CONTENT = prs.slide_layouts[0], prs.slide_layouts[1]


def runs(p, text, size, color, bold=False, italic=False, font='Calibri'):
    """Add text to paragraph p; **bold** segments are bold."""
    for k, part in enumerate(re.split(r'\*\*', text)):
        if not part:
            continue
        r = p.add_run();  r.text = part
        f = r.font;  f.size = Pt(size);  f.bold = bold or (k % 2 == 1);  f.italic = italic
        f.color.rgb = rgb(color);  f.name = font


def textbox(s, x, y, w, h, paras, size=12.5, color=TEXT, bold=False, italic=False, align=PP_ALIGN.LEFT, space=4):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame;  tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, t in enumerate(paras if isinstance(paras, list) else [paras]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align;  p.space_after = Pt(space)
        runs(p, t, size, color, bold, italic)
    return tb


def footer(s, n):
    """Date, label and slide number in the template's footer bar (cloned from the layout)."""
    for ph in L_CONTENT.placeholders:
        idx = ph.placeholder_format.idx
        if idx in (12, 13, 14):
            el = copy.deepcopy(ph._element);  s.shapes._spTree.append(el)
            sh = s.shapes[-1];  sh.text_frame.text = {12: 'CONFIDENTIAL', 13: str(n), 14: DATE}[idx]


def slide(title, subtitle=None):
    s = prs.slides.add_slide(L_CONTENT)
    s.shapes.title.text = title
    for r in s.shapes.title.text_frame.paragraphs[0].runs:
        r.font.size = Pt(28)
    if subtitle:
        textbox(s, 0.1, 0.7, 10.75, 0.36, subtitle, size=15, color=TERRA, italic=True)
    footer(s, len(prs.slides))
    return s


def card(s, x, y, w, h, header=None, body=(), hcolor=BLUE, size=14):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.10                                 # as the template cards
    sh.fill.solid();  sh.fill.fore_color.rgb = rgb(CARD);  sh.line.fill.background();  sh.shadow.inherit = False
    top = y + 0.1
    if header:
        textbox(s, x + 0.2, top, w - 0.4, 0.38, header, size=17, color=hcolor, bold=True);  top += 0.47
    if body:
        textbox(s, x + 0.2, top, w - 0.4, h - (top - y) - 0.1, list(body), size=size)
    return sh


def figure(s, path, x, y, w, h):
    """Picture fitted into the box (aspect kept, centred)."""
    iw, ih = Image.open(path).size
    sc = min(w / iw, h / ih);  pw, ph = iw * sc, ih * sc
    s.shapes.add_picture(path, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def idea(s, x, y, w, h, text):
    """Dashed box: a figure still to be decided (what it could show, and where it exists)."""
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.04
    sh.fill.background();  sh.line.color.rgb = rgb('9A998F');  sh.line.width = Pt(1.25)
    sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH;  sh.shadow.inherit = False
    tf = sh.text_frame;  tf.word_wrap = True;  tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.25)
    for i, t in enumerate(['FIGURE TO DECIDE'] + text):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER;  p.space_after = Pt(4)
        runs(p, t, 12 if i else 11, MUTED, bold=(i == 0), italic=(i > 0))


paper = lambda n: os.path.join(A, 'paper', n + '.png')
asset = lambda n: os.path.join(A, n + '.png')

# ================================================================== 1. title
s = prs.slides.add_slide(L_TITLE)
s.shapes.title.text = 'From channel charting to EMS configuration'
s.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(26)
for ph in s.placeholders:
    if ph.placeholder_format.idx == 10:                    # the layout's subtitle box is not shown:
        ph._element.getparent().remove(ph._element)        # replaced by a plain text box
    elif ph.placeholder_format.idx == 11:
        ph.text = 'M. Maleki, R. Aghazadeh Ayoubi, M. Mizmizi, U. Spagnolini  |  Politecnico di Milano  |  ' + DATE
        ph.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
textbox(s, 2.35, 6.02, 8.62, 0.4, 'AoA estimation and beam selection with few sweeps', size=17, color='FFFFFF', italic=True, align=PP_ALIGN.CENTER)

# ================================================================== PART A: the paper
s = slide('Smart radio environment: two EMSs help one base station',
          'Our TWC paper: electromagnetic surfaces (EMSs) improve channel charting for localisation.')
figure(s, paper('system_model'), 0.45, 1.2, 6.9, 5.55)
card(s, 7.6, 1.2, 5.2, 2.35, 'Scenario', [
    '30 GHz urban street, Sionna ray tracing; BS with 32 antennas.',
    'Two passive EMSs (60x60 elements, lambda/2) on building walls; **3200 users** on a grid.',
    'A user reaches the BS directly (h^d) and through each EMS (h^EMS).'])
card(s, 7.6, 3.7, 5.2, 3.05, 'Channel charting', [
    'Each user: covariance of its uplink channel at the BS.',
    'Log-Euclidean distance between covariances -> **semi-supervised t-SNE** (15% labelled anchors).',
    'The chart places users so that similar channels are neighbours: with position labels it estimates **where the user is**.'], hcolor=TERRA)

s = slide('Paper result: EMS configurations make the chart more accurate',
          'A few static EMS configurations (slots) give one chart each; fusing them per user halves the error.')
figure(s, paper('cdf_ms_st'), 0.45, 1.2, 6.3, 5.55)
card(s, 7.0, 1.2, 5.8, 2.6, 'Multi-slot acquisition', [
    'K static EMS phase profiles, one uplink acquisition each, **no knowledge of the user needed**.',
    'One chart per configuration; per user, charts weighted by their quality (trustworthiness / continuity).'])
card(s, 7.0, 3.95, 5.8, 2.8, 'Positioning error (mean)', [
    'No EMS: **17.5 m**  ->  best single configuration: **11 m**',
    'Per-acquisition weighting: **8.7 m**  ->  point-wise weighting: **7.2 m**',
    'Message: the EMS shapes the radio map, and a few configurations are enough.'], hcolor=TERRA)

# ================================================================== PART B: new goal
s = slide('New goal: configure the EMS itself',
          'To steer the EMS toward a user we need its direction; to learn the direction we must first illuminate it.')
card(s, 0.55, 1.2, 6.2, 2.55, 'A reconfigurable surface needs a direction', [
    'The EMS (a RIS) reflects toward one direction per phase profile: a **codebook** of beams.',
    'Best beam for a user = the one pointing at its **angle of arrival (AoA)** at the EMS, which is unknown.'])
card(s, 0.55, 3.9, 6.2, 2.85, 'The cost of finding it by search', [
    'Exhaustive sweep: try every beam, keep the strongest: **68 beams** (60x60), **272** (240x240).',
    'Guided searches (greedy, hierarchical) need fewer sweeps but still several per user.',
    'Cyclic dilemma: steering needs the angle, measuring needs steering.'], hcolor=TERRA)
figure(s, paper('asilomar_full_sweep'), 7.1, 1.2, 5.7, 5.55)

s = slide('Proposed: channel charting of the angle, then a look-up table',
          'The paper\'s static slots become training acquisitions; the chart learns the AoA instead of the position.')
figure(s, paper('asilomar_system'), 0.45, 1.15, 12.4, 3.5)
card(s, 0.55, 4.8, 4.0, 1.95, '1  K training acquisitions', [
    'Static EMS configurations, as in the paper. **Informed**: chosen from the angular spread of the anchors.'])
card(s, 4.67, 4.8, 4.0, 1.95, '2  Angular chart', [
    'Semi-supervised t-SNE labelled with the **azimuth at the EMS**; per user, the most trustworthy of the K charts.'], hcolor=TERRA)
card(s, 8.79, 4.8, 4.0, 1.95, '3  Look-up table', [
    'Azimuth -> codeword through the steering law: the EMS serves the user **without sweeping**.'])

s = slide('Who needs the EMS? The users in outage',
          'Only users without a usable BS link gain from the EMS: in our scenario, 28% of the users.')
figure(s, asset('scenario_link_classes'), 0.45, 1.15, 5.6, 5.65)
card(s, 6.3, 1.2, 6.5, 1.55, 'Line of sight: 811 users (25%)', [
    'Strong direct link: the EMS is not needed.'])
card(s, 6.3, 2.9, 6.5, 1.75, 'Reflection only: 1504 users (47%)', [
    'A wall reflection reaches the BS, about **30 dB stronger** than the best EMS path: the EMS adds less than 0.01 dB.'], hcolor=TERRA)
card(s, 6.3, 4.8, 6.5, 1.95, 'Full NLOS: 885 users (28%)  ->  our target', [
    'No BS-UE channel at all: the EMS is their **only link**. All results below are for these users.'], hcolor=BLUE)

# ================================================================== PART C: results, one EMS at a time
s = slide('Evaluation setup', 'Each EMS evaluated on its own: its chart, its beams, the users it reaches.')
card(s, 0.55, 1.2, 6.05, 2.6, 'Surfaces and codebooks', [
    'Real 60x60 EMS (ray-traced); synthetic 120x120 ... 240x240 (fitted spherical-wave model).',
    'Codebook from the beamwidth: **68 beams** (60x60) to **272** (240x240).',
    'Estimated channels: 23 dBm, -92 dBm noise, 60 dB pilot gain.'])
card(s, 6.75, 1.2, 6.05, 2.6, 'Training (K acquisitions)', [
    '**Informed**: training beams cover the angular groups of the anchors (prior).',
    '**Coarse-to-fine**: one fixed ever-finer order (no prior).',
    '**Random**: random codewords (no prior), 10 draws.'], hcolor=TERRA)
card(s, 0.55, 3.95, 6.05, 2.8, 'Baselines (same sweep budget)', [
    '**Greedy search** and **hierarchical search** over the beams pointing at the user groups (same prior).',
    'Exhaustive search (68 ... 272 sweeps) as the upper bound.'], hcolor=TERRA)
card(s, 6.75, 3.95, 6.05, 2.8, 'Metrics', [
    '**Pointing error** of the chosen beam, in beamwidths (1 beamwidth = 1.7 deg at broadside, 60x60).',
    '**SNR gain** of the chosen beam vs the unsteered EMS (zero phase), median over users.'])

s = slide('Case 1: EMS1 installed', 'Which users in outage can EMS1 serve, and how well does the chart locate them?')
figure(s, asset('coverage_EMS1'), 0.45, 1.15, 5.9, 5.65)
idea(s, 6.6, 1.2, 6.2, 5.55, ['Per-user error map of the t-SNE angle estimate, K = 3 (EMS1 trained alone).',
                               'Exists: serving_deployments/.../figures/per_user/1_informed_prior/map_M60_raytraced_K3',
                               'or: estimated vs true azimuth (truevsest_M60_raytraced).'])

s = slide('Case 1: results', 'With 1-3 sweeps the chart beats a guided search; with more sweeps, search catches up.')
idea(s, 0.55, 1.2, 8.2, 5.55, ['Median pointing error and median SNR gain vs number of sweeps (EMS1 trained alone):',
                                't-SNE, t-SNE + neighbours, greedy and hierarchical search (clusters), exhaustive.',
                                'Exists as one column of figures_previous_style_newmetric/1_informed_prior; I can draw a 2-panel version.'])
card(s, 9.0, 1.2, 3.8, 5.55, 'Key numbers (real 60x60)', [
    'SNR gain at **1 sweep**: t-SNE **16.3 dB**, greedy 9.5 dB, hierarchical 8.7 dB.',
    'Angle saturates after **3 acquisitions** (median 2.1 deg).',
    'From about 5 sweeps the guided search reaches t-SNE.'], hcolor=TERRA)

s = slide('Case 2: EMS2 installed', 'EMS2 sees its users in three separate angular groups: one training beam lights only one.')
figure(s, asset('coverage_EMS2'), 0.45, 1.15, 5.9, 5.65)
idea(s, 6.6, 1.2, 6.2, 5.55, ['Estimated vs true azimuth for K = 1 and K = 3 (EMS2): with one acquisition, whole groups are misplaced; with three, they line up.',
                               'Exists: figures/per_user/1_informed_prior/truevsest_M60_raytraced'])

s = slide('Case 2: results', 'Real 60x60: a guided search is as good early on; on larger surfaces the chart leads.')
idea(s, 0.55, 1.2, 6.05, 5.55, ['Median SNR gain vs sweeps, EMS2 trained alone, real 60x60.'])
idea(s, 6.75, 1.2, 6.05, 5.55, ['Same for the 240x240 surface: t-SNE leads at 2-5 sweeps (25.6 vs 17.0 dB at 3 sweeps).'])

s = slide('Training choice matters', 'Placing the training beams on the user groups (informed) beats random or fixed training.')
idea(s, 0.55, 1.2, 8.2, 5.55, ['Median angle error vs K: informed vs coarse-to-fine vs random, EMS1 and EMS2 (trained alone).',
                                'Numbers exist (EMS2, real 60x60): K = 3: 3.7 deg informed vs 37 deg random; K = 10: 1.9 vs 14 deg.'])
card(s, 9.0, 1.2, 3.8, 5.55, 'Why', [
    'A chart places a user well only if a training beam lights it.',
    'Informed training needs **one beam per user group**: about 3 for EMS2, instead of 6 (coarse-to-fine) or 10+ (random).'], hcolor=TERRA)

s = slide('Takeaways')
card(s, 0.55, 1.2, 12.25, 5.55, None, [
    '**From localisation to configuration**: the same static EMS slots, labelled with the AoA, let the BS choose the EMS beam without sweeping.',
    '**Target users**: only full-NLOS users gain from the EMS (28% here); line-of-sight and wall-reflection users do not need it.',
    '**Few sweeps**: with 1-3 acquisitions the chart serves users better than a guided search (EMS1; larger surfaces), and saturates at about 3.',
    '**Informed training** (one beam per user group) is the key ingredient; a few users still get a wrong group (the tail).',
    'Next: hybrid t-SNE + short search for the tail; joint use of both EMSs (backup).'], size=15)

# ================================================================== PART D (backup)
s = slide('Backup: both EMSs present', 'Joint channel, separate inference per EMS; or one joint chart for both angles.')
idea(s, 0.55, 1.2, 6.05, 5.55, ['Joint channel, separate charts: EMS2 at K = 1 improves from 46 deg (alone) to 9 deg (joint).',
                                 'Exists: serving_single_ems/.../figures_TEMP/fig_angle_joint_vs_separated'])
idea(s, 6.75, 1.2, 6.05, 5.55, ['One joint 2-D chart (both angles): clearly worse (6-15 deg vs 2-4 deg).',
                                 'Exists: serving_joint_chart/.../figures/shell_order_bothReached'])

prs.save(OUT)
print('saved', OUT, len(prs.slides), 'slides')
